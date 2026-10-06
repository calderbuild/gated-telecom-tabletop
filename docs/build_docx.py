"""Build docs/out/report.docx from docs/report.md: pandoc, then compact styling (the official
template's styles break pandoc tables, so headings and structure follow the template, not its styles)."""
import subprocess
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).parent
OUT = HERE / "out" / "report.docx"
OUT.parent.mkdir(exist_ok=True)
subprocess.run(["pandoc", str(HERE / "report.md"), "--resource-path", str(HERE), "-o", str(OUT)], check=True)
d = docx.Document(OUT)
for sec in d.sections:
    sec.left_margin = sec.right_margin = Inches(0.7)
    sec.top_margin = sec.bottom_margin = Inches(0.6)
for name, size in [("Normal", 10), ("Body Text", 10), ("First Paragraph", 10), ("Compact", 10), ("Verbatim Char", 8.5),
                   ("Title", 14), ("Heading 2", 11.5)]:
    if name in [s.name for s in d.styles]:
        st = d.styles[name]
        st.font.size = Pt(size)
        if not hasattr(st, "paragraph_format"):
            continue
        st.paragraph_format.space_after = Pt(3)
        st.paragraph_format.space_before = Pt(6 if name.startswith("Heading") else 0)
for t in d.tables:
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "999999")
        b.append(e)
    t._tbl.tblPr.append(b)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(8)
d.save(OUT)
